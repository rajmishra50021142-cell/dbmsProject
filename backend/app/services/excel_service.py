"""
Service for Excel workbook ingestion, inspection, heuristic suggestions,
and verified normalized spreadsheet export.
"""

import io
import re
from typing import List, Dict, Any, Optional, Set, Tuple
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.schemas.domain_contracts import (
    CanonicalSchemaInput,
    FunctionalDependency,
    MultivaluedDependency,
)
from app.schemas.excel_schemas import (
    ExcelInspectResponse,
    ExcelIssue,
    ExcelSuggestResponse,
    ObservedKey,
    ObservedFD,
    AtomicityReviewItem,
)
from app.normalization.normalization_engine import analyze_full_normalization
from app.normalization.candidate_key_engine import find_all_candidate_keys
from app.decomposition.decomposition_engine import analyze_decomposition_pipeline


class ExcelService:
    """Handles Excel workbook parsing, validation, suggestion mining, and export."""

    @staticmethod
    def inspect_workbook(
        file_bytes: bytes,
        filename: str,
        sheet_name: Optional[str] = None,
        max_preview_rows: int = 50,
    ) -> ExcelInspectResponse:
        """
        Safely inspects an uploaded Excel (.xlsx) file without executing macros.
        Extracts available sheets, column headers, preview records, and flags issues.
        """
        # Validate magic bytes for ZIP/OOXML format (.xlsx)
        if not file_bytes.startswith(b"PK\x03\x04"):
            raise ValueError("Invalid file format. Uploaded file is not a valid Excel (.xlsx) workbook.")

        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
        sheet_names = wb.sheetnames

        if not sheet_names:
            wb.close()
            raise ValueError("The uploaded Excel workbook contains no worksheets.")

        active_sheet_name = sheet_name if sheet_name and sheet_name in sheet_names else sheet_names[0]
        sheet = wb[active_sheet_name]

        issues: List[ExcelIssue] = []
        raw_rows = list(sheet.iter_rows(values_only=True))
        wb.close()

        if not raw_rows:
            return ExcelInspectResponse(
                filename=filename,
                sheets=sheet_names,
                selected_sheet=active_sheet_name,
                headers=[],
                total_rows=0,
                preview_rows=[],
                issues=[
                    ExcelIssue(
                        severity="warning",
                        code="EMPTY_WORKSHEET",
                        message=f"Worksheet '{active_sheet_name}' is completely empty.",
                    )
                ],
            )

        # Locate first non-empty row as header
        header_row_index = -1
        raw_headers = []
        for idx, row in enumerate(raw_rows):
            if any(cell is not None and str(cell).strip() != "" for cell in row):
                header_row_index = idx
                raw_headers = [cell for cell in row]
                break

        if header_row_index == -1:
            return ExcelInspectResponse(
                filename=filename,
                sheets=sheet_names,
                selected_sheet=active_sheet_name,
                headers=[],
                total_rows=0,
                preview_rows=[],
                issues=[
                    ExcelIssue(
                        severity="error",
                        code="NO_HEADER_ROW",
                        message="No header row with attribute names could be identified.",
                    )
                ],
            )

        # Normalize headers and detect blank or duplicate column names
        headers: List[str] = []
        seen_headers: Set[str] = set()

        for col_idx, cell in enumerate(raw_headers):
            col_letter = get_column_letter(col_idx + 1)
            if cell is None or str(cell).strip() == "":
                gen_name = f"Column_{col_letter}"
                headers.append(gen_name)
                issues.append(
                    ExcelIssue(
                        severity="warning",
                        code="EMPTY_COLUMN_HEADER",
                        message=f"Empty header in column {col_letter}; auto-assigned name '{gen_name}'.",
                        column=gen_name,
                        row=header_row_index + 1,
                    )
                )
            else:
                h_name = str(cell).strip()
                if h_name in seen_headers:
                    # Rename duplicate column
                    dup_name = f"{h_name}_{col_letter}"
                    headers.append(dup_name)
                    issues.append(
                        ExcelIssue(
                            severity="warning",
                            code="DUPLICATE_HEADER",
                            message=f"Duplicate column header '{h_name}' at column {col_letter}; renamed to '{dup_name}'.",
                            column=dup_name,
                            row=header_row_index + 1,
                        )
                    )
                else:
                    headers.append(h_name)
                    seen_headers.add(h_name)

        # Extract data rows
        data_rows_raw = raw_rows[header_row_index + 1 :]
        parsed_rows: List[Dict[str, Any]] = []

        for row_idx, r in enumerate(data_rows_raw):
            # Check if entire row is empty
            if all(cell is None or str(cell).strip() == "" for cell in r):
                continue

            row_dict: Dict[str, Any] = {}
            for col_idx, h in enumerate(headers):
                val = r[col_idx] if col_idx < len(r) else None
                # Standardize strings, preserve numbers
                if isinstance(val, str):
                    val = val.strip()
                row_dict[h] = val
            parsed_rows.append(row_dict)

        total_rows = len(parsed_rows)
        preview = parsed_rows[:max_preview_rows]

        # Scan for atomicity candidates in preview
        for row_i, row in enumerate(preview):
            for h in headers:
                val = row.get(h)
                if isinstance(val, str) and ("," in val or ";" in val or "|" in val):
                    # Check if it looks like a list
                    parts = [p.strip() for p in re.split(r"[,;|]", val) if p.strip()]
                    if len(parts) > 1 and not re.match(r"^\d+([,.]\d+)?$", val):
                        issues.append(
                            ExcelIssue(
                                severity="info",
                                code="POTENTIAL_NON_ATOMIC",
                                message=f"Column '{h}' in row {row_i + 1} contains delimited values ('{val}'). Review for 1NF atomicity.",
                                column=h,
                                row=row_i + 2,
                            )
                        )
                        break

        return ExcelInspectResponse(
            filename=filename,
            sheets=sheet_names,
            selected_sheet=active_sheet_name,
            headers=headers,
            total_rows=total_rows,
            preview_rows=preview,
            issues=issues[:20],  # Cap issues to 20 for readability
        )

    @staticmethod
    def suggest_schema_and_dependencies(
        attributes: List[str],
        rows: List[Dict[str, Any]],
        relation_name: Optional[str] = "UPLOADED_RELATION",
    ) -> ExcelSuggestResponse:
        """
        Mode B: Heuristic observation mining on uploaded spreadsheet records.
        Suggests unique keys, potential functional dependencies, and atomicity review items.
        Clearly notes that observed sample data patterns are not formal constraints.
        """
        rel_name = relation_name or "UPLOADED_RELATION"
        suggested_keys: List[ObservedKey] = []
        suggested_fds: List[ObservedFD] = []
        atomicity_reviews: List[AtomicityReviewItem] = []
        data_quality_notes: List[str] = []

        total_rows = len(rows)
        if total_rows == 0:
            return ExcelSuggestResponse(
                relation_name=rel_name,
                suggested_keys=[],
                suggested_fds=[],
                atomicity_reviews=[],
                data_quality_notes=["Spreadsheet contains no data rows to analyze."],
            )

        # 1. Uniqueness / Candidate Key Suggestions (Single and Pairs)
        single_unique: Set[str] = set()
        for attr in attributes:
            values = [r.get(attr) for r in rows if r.get(attr) is not None and str(r.get(attr)).strip() != ""]
            if len(values) == total_rows and len(set(values)) == total_rows:
                single_unique.add(attr)
                suggested_keys.append(
                    ObservedKey(
                        attributes=[attr],
                        is_unique=True,
                        explanation=f"Attribute '{attr}' contains strictly unique values across all {total_rows} uploaded rows.",
                    )
                )

        # If no single unique key found, inspect 2-attribute pairs
        if not suggested_keys and len(attributes) >= 2:
            for i in range(len(attributes)):
                for j in range(i + 1, len(attributes)):
                    a1, a2 = attributes[i], attributes[j]
                    pair_values = [
                        (r.get(a1), r.get(a2))
                        for r in rows
                        if r.get(a1) is not None and r.get(a2) is not None
                    ]
                    if len(pair_values) == total_rows and len(set(pair_values)) == total_rows:
                        suggested_keys.append(
                            ObservedKey(
                                attributes=[a1, a2],
                                is_unique=True,
                                explanation=f"Composite key {{{a1}, {a2}}} is unique across all {total_rows} observed rows.",
                            )
                        )

        # 2. Heuristic Functional Dependency Mining
        # Check X -> Y where X != Y and X is not already a unique key
        candidate_dets: List[List[str]] = [[a] for a in attributes if a not in single_unique]
        # Also include pairs if single attribute determinants are few
        if len(candidate_dets) < 5 and len(attributes) >= 3:
            for i in range(len(attributes)):
                for j in range(i + 1, len(attributes)):
                    if attributes[i] not in single_unique and attributes[j] not in single_unique:
                        candidate_dets.append([attributes[i], attributes[j]])

        for det in candidate_dets:
            for target in attributes:
                if target in det:
                    continue

                # Check if for every row, det uniquely determines target
                mapping: Dict[Tuple, Any] = {}
                holds = True
                for r in rows:
                    key_val = tuple(r.get(k) for k in det)
                    t_val = r.get(target)
                    # Ignore null determinants
                    if any(kv is None or str(kv).strip() == "" for kv in key_val):
                        continue
                    if key_val in mapping:
                        if mapping[key_val] != t_val:
                            holds = False
                            break
                    else:
                        mapping[key_val] = t_val

                # If holds and has non-trivial grouping (more than 1 distinct determinant value)
                if holds and len(mapping) > 1 and len(mapping) < total_rows:
                    det_str = ", ".join(det)
                    suggested_fds.append(
                        ObservedFD(
                            left=det,
                            right=[target],
                            holds=True,
                            explanation=f"In all {total_rows} sample rows, identical values of {{{det_str}}} always map to the exact same value of '{target}'.",
                            is_heuristic=True,
                        )
                    )

        # 3. Atomicity Reviews (check for multiple values in cells)
        for attr in attributes:
            found_count = 0
            for idx, r in enumerate(rows):
                val = r.get(attr)
                if isinstance(val, str) and ("," in val or ";" in val or "|" in val):
                    # Check if it has multiple tokens
                    parts = [p.strip() for p in re.split(r"[,;|]", val) if p.strip()]
                    if len(parts) > 1 and not re.match(r"^\d+([,.]\d+)?$", val):
                        delim = "," if "," in val else (";" if ";" in val else "|")
                        atomicity_reviews.append(
                            AtomicityReviewItem(
                                attribute=attr,
                                row_index=idx + 1,
                                sample_value=val,
                                delimiter_detected=delim,
                                message=f"Contains multiple values separated by '{delim}'. If these represent multiple domain entries, normalize to 1NF by splitting into separate rows or relations.",
                            )
                        )
                        found_count += 1
                        if found_count >= 2:
                            break

        # 4. Data Quality Notes
        # Check duplicate rows
        row_tuples = [tuple(r.get(a) for a in attributes) for r in rows]
        dup_count = len(row_tuples) - len(set(row_tuples))
        if dup_count > 0:
            data_quality_notes.append(
                f"Found {dup_count} duplicate row(s) in the uploaded worksheet. In relational databases, relations are mathematical sets and should not contain duplicate tuples."
            )

        # Check null values
        null_counts = {
            a: sum(1 for r in rows if r.get(a) is None or str(r.get(a)).strip() == "")
            for a in attributes
        }
        high_nulls = [a for a, c in null_counts.items() if c > 0]
        if high_nulls:
            null_summary = ", ".join([f"{a} ({null_counts[a]} nulls)" for a in high_nulls[:3]])
            data_quality_notes.append(f"Missing/blank cells detected in: {null_summary}.")

        return ExcelSuggestResponse(
            relation_name=rel_name,
            suggested_keys=suggested_keys[:5],
            suggested_fds=suggested_fds[:10],
            atomicity_reviews=atomicity_reviews[:10],
            data_quality_notes=data_quality_notes,
        )

    @staticmethod
    def validate_rows_against_dependencies(
        attributes: List[str],
        functional_dependencies: List[FunctionalDependency],
        rows: List[Dict[str, Any]],
    ) -> List[Dict[str, Any]]:
        """
        Validates whether uploaded spreadsheet records satisfy the declared functional dependencies.
        Returns a list of specific contradiction violations.
        """
        contradictions = []

        for fd in functional_dependencies:
            left_attrs = fd.left
            right_attrs = fd.right

            # Only check if all attributes exist in rows
            if not all(a in attributes for a in left_attrs + right_attrs):
                continue

            value_map: Dict[Tuple, Tuple[int, Dict[str, Any]]] = {}
            for r_idx, row in enumerate(rows):
                left_val = tuple(row.get(a) for a in left_attrs)
                # Skip if any LHS is null
                if any(v is None or str(v).strip() == "" for v in left_val):
                    continue

                right_val = {a: row.get(a) for a in right_attrs}

                if left_val in value_map:
                    orig_row_idx, orig_right_val = value_map[left_val]
                    if orig_right_val != right_val:
                        left_str = ", ".join(f"{a}={row.get(a)}" for a in left_attrs)
                        contradictions.append({
                            "dependency": f"{', '.join(left_attrs)} → {', '.join(right_attrs)}",
                            "message": (
                                f"Row {orig_row_idx + 1} and Row {r_idx + 1} share identical determinant ({left_str}), "
                                f"but have differing dependent values ({orig_right_val} vs {right_val})."
                            ),
                            "row_1": orig_row_idx + 1,
                            "row_2": r_idx + 1,
                        })
                        break
                else:
                    value_map[left_val] = (r_idx, right_val)

        return contradictions

    @staticmethod
    def generate_normalized_workbook(
        relation_name: str,
        attributes: List[str],
        functional_dependencies: List[FunctionalDependency],
        multivalued_dependencies: List[MultivaluedDependency],
        candidate_keys: List[List[str]],
        rows: List[Dict[str, Any]],
        target_normal_form: str = "3NF",
    ) -> bytes:
        """
        Executes real normalization decomposition via Phase 7 decomposition engine and Phase 5/6
        analysis engine, projects uploaded spreadsheet tuples onto the resulting relations, removes
        duplicate projected tuples, and generates a formatted Excel workbook with individual sheets
        and a 'Normalization Summary' sheet.
        """
        # 1. Run full normalization analysis for formal verification
        canonical_input = CanonicalSchemaInput(
            name=relation_name,
            attributes=attributes,
            candidate_keys=candidate_keys,
            functional_dependencies=functional_dependencies,
            multivalued_dependencies=multivalued_dependencies,
        )
        full_analysis = analyze_full_normalization(canonical_input)

        # 2. Run formal decomposition pipeline
        target_nf = target_normal_form.upper()
        if target_nf not in ["2NF", "3NF", "4NF"]:
            target_nf = "3NF"

        decomp_res = analyze_decomposition_pipeline(
            relation_name=relation_name,
            attributes=attributes,
            fds=functional_dependencies,
            mvds=multivalued_dependencies,
            target_normal_form=target_nf,
            candidate_keys=candidate_keys,
        )

        final_relations = decomp_res.final_relations
        if not final_relations:
            # Fallback to source relation if already in target normal form
            from app.schemas.domain_contracts import DecomposedRelation
            final_relations = [
                DecomposedRelation(
                    name=relation_name,
                    attributes=attributes,
                    candidate_keys=candidate_keys or find_all_candidate_keys(attributes, functional_dependencies),
                    primary_key=candidate_keys[0] if candidate_keys else attributes[:1],
                    purpose=f"Relation already satisfies {target_nf}; no decomposition required.",
                    source_relation=relation_name,
                )
            ]

        # 3. Create Excel workbook with openpyxl
        wb = openpyxl.Workbook()
        # Remove default sheet
        wb.remove(wb.active)

        # Styling constants
        header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")  # Indigo
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        title_font = Font(name="Calibri", size=14, bold=True, color="1E1B4B")
        sub_title_font = Font(name="Calibri", size=11, bold=True, color="4338CA")
        bold_font = Font(name="Calibri", size=11, bold=True, color="1E293B")
        regular_font = Font(name="Calibri", size=11, color="334155")
        thin_border = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1"),
        )
        accent_fill = PatternFill(start_color="EEF2FF", end_color="EEF2FF", fill_type="solid")
        zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

        # 4. Generate a worksheet for each decomposed relation
        used_sheet_names: Set[str] = set()

        for rel in final_relations:
            # Sanitize sheet name (Excel limits: max 31 chars, no []:*?/\)
            raw_title = rel.name or "SubRelation"
            clean_title = re.sub(r"[\[\]:*?/\\]", "_", raw_title)[:28]
            sheet_title = clean_title
            counter = 1
            while sheet_title in used_sheet_names:
                sheet_title = f"{clean_title}_{counter}"[:31]
                counter += 1
            used_sheet_names.add(sheet_title)

            ws = wb.create_sheet(title=sheet_title)
            ws.views.sheetView[0].showGridLines = True

            # Write Relation Header Banner
            ws.append([f"Relation: {rel.name}"])
            ws.cell(row=1, column=1).font = title_font

            pk_str = ", ".join(rel.primary_key) if rel.primary_key else "None declared"
            ws.append([f"Primary Key: {{{pk_str}}} | Purpose: {rel.purpose or 'Decomposed sub-relation'}"])
            ws.cell(row=2, column=1).font = regular_font
            ws.append([])  # blank row

            # Table Column Headers
            ws.append(rel.attributes)
            for col_idx in range(1, len(rel.attributes) + 1):
                cell = ws.cell(row=4, column=col_idx)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = thin_border

            # Project rows and eliminate duplicates (relational projection)
            seen_tuples: Set[Tuple] = set()
            row_count = 0

            for orig_row in rows:
                proj_tuple = tuple(orig_row.get(attr) for attr in rel.attributes)
                if proj_tuple not in seen_tuples:
                    seen_tuples.add(proj_tuple)
                    row_cells = list(proj_tuple)
                    ws.append(row_cells)
                    row_count += 1
                    current_row = 4 + row_count
                    for c_idx in range(1, len(rel.attributes) + 1):
                        c = ws.cell(row=current_row, column=c_idx)
                        c.font = regular_font
                        c.border = thin_border
                        if row_count % 2 == 0:
                            c.fill = zebra_fill

            # Auto-fit column widths
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    if cell.row >= 4 and cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                ws.column_dimensions[col_letter].width = max(max_len + 5, 12)

        # 5. Generate Normalization Summary Worksheet
        ws_summary = wb.create_sheet(title="Normalization Summary")
        ws_summary.views.sheetView[0].showGridLines = True

        # Header Title
        ws_summary.append(["NORMALIZATION LAB — TRANSFORMATION SUMMARY"])
        ws_summary.cell(row=1, column=1).font = title_font
        ws_summary.append(["Generated by Normalization Lab — Formal Relational Normalization & Decomposition Engine"])
        ws_summary.cell(row=2, column=1).font = regular_font
        ws_summary.append([])

        # Section 1: Original Schema
        ws_summary.append(["1. ORIGINAL RELATION SCHEMA"])
        ws_summary.cell(row=4, column=1).font = sub_title_font
        ws_summary.append(["Relation Name", relation_name])
        ws_summary.append(["Original Attributes", ", ".join(attributes)])

        fd_strings = [f"{', '.join(fd.left)} → {', '.join(fd.right)}" for fd in functional_dependencies]
        ws_summary.append(["Declared Functional Dependencies", "; ".join(fd_strings) or "None declared"])

        mvd_strings = [f"{', '.join(mvd.left)} ↠ {', '.join(mvd.right)}" for mvd in multivalued_dependencies]
        ws_summary.append(["Declared Multivalued Dependencies", "; ".join(mvd_strings) or "None declared"])

        keys_list = [f"{{{', '.join(k)}}}" for k in full_analysis.candidate_keys]
        ws_summary.append(["Confirmed Candidate Keys", ", ".join(keys_list) or "None computed"])
        ws_summary.append(["Highest Confirmed Normal Form", full_analysis.highest_confirmed_normal_form])
        ws_summary.append([])

        # Section 2: Normal Form Findings
        ws_summary.append(["2. FORMAL NORMAL FORM STATUSES"])
        ws_summary.cell(row=12, column=1).font = sub_title_font
        ws_summary.append(["Normal Form", "Status", "Violations Count", "Academic Notes"])
        for c_idx in range(1, 5):
            c = ws_summary.cell(row=13, column=c_idx)
            c.fill = header_fill
            c.font = header_font
            c.border = thin_border

        nf_rows = [
            ("1NF (First Normal Form)", full_analysis.nf1.status, str(len(full_analysis.nf1.violations)), full_analysis.nf1.message or "Atomic values"),
            ("2NF (Second Normal Form)", full_analysis.nf2.status, str(len(full_analysis.nf2.partial_dependencies or [])), full_analysis.nf2.message or "No partial dependencies"),
            ("3NF (Third Normal Form)", full_analysis.nf3.status if full_analysis.nf3 else "N/A", str(len(full_analysis.nf3.violations or [])) if full_analysis.nf3 else "0", full_analysis.nf3.message or "" if full_analysis.nf3 else "N/A"),
            ("4NF (Fourth Normal Form)", full_analysis.nf4.status if full_analysis.nf4 else "N/A", str(len(full_analysis.nf4.violations or [])) if full_analysis.nf4 else "0", full_analysis.nf4.message or "" if full_analysis.nf4 else "N/A"),
        ]

        curr_r = 14
        for nf_name, st, v_cnt, notes in nf_rows:
            ws_summary.append([nf_name, st, v_cnt, notes])
            for c_idx in range(1, 5):
                cell = ws_summary.cell(row=curr_r, column=c_idx)
                cell.font = regular_font
                cell.border = thin_border
            curr_r += 1

        ws_summary.append([])
        curr_r += 1

        # Section 3: Resulting Relations Summary
        ws_summary.append(["3. DECOMPOSED RELATIONS GENERATED"])
        ws_summary.cell(row=curr_r, column=1).font = sub_title_font
        curr_r += 1

        ws_summary.append(["Relation Name", "Attributes", "Primary Key", "Worksheet Name", "Purpose"])
        for c_idx in range(1, 6):
            cell = ws_summary.cell(row=curr_r, column=c_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.border = thin_border
        curr_r += 1

        for rel in final_relations:
            pk = ", ".join(rel.primary_key) if rel.primary_key else "None"
            ws_summary.append([rel.name, ", ".join(rel.attributes), f"{{{pk}}}", rel.name[:28], rel.purpose or "Decomposed relation"])
            for c_idx in range(1, 6):
                cell = ws_summary.cell(row=curr_r, column=c_idx)
                cell.font = regular_font
                cell.border = thin_border
            curr_r += 1

        ws_summary.append([])
        curr_r += 1

        # Section 4: Formal Verification (Lossless Join & Dependency Preservation)
        ws_summary.append(["4. FORMAL MATHEMATICAL VERIFICATION"])
        ws_summary.cell(row=curr_r, column=1).font = sub_title_font
        curr_r += 1

        # Check if decomposition plans exist with verification details
        is_lossless = False
        lossless_method = "Tableau Chase Algorithm"
        is_preserved = True
        preserved_summary = "All functional dependencies preserved."

        if decomp_res.plans:
            last_plan = decomp_res.plans[-1]
            is_lossless = last_plan.lossless_join.is_lossless
            lossless_method = last_plan.lossless_join.method
            is_preserved = last_plan.dependency_preservation.is_preserved
            preserved_summary = (
                f"{len(last_plan.dependency_preservation.preserved_dependencies)} FDs preserved, "
                f"{len(last_plan.dependency_preservation.non_preserved_dependencies)} FDs non-preserved"
            )
        else:
            is_lossless = True
            lossless_method = "Identity mapping (no decomposition needed)"

        ws_summary.append(["Lossless-Join Verification", "VERIFIED LOSSLESS" if is_lossless else "UNVERIFIED / LOSSY", f"Method: {lossless_method}"])
        ws_summary.append(["Dependency Preservation", "VERIFIED PRESERVED" if is_preserved else "DEPENDENCIES LOST", preserved_summary])
        ws_summary.append(["Verification Engine", "Relational Algorithm Matrix (Phases 3–7)", "Tableau Chase & Polynomial-Time Closure Projection"])
        ws_summary.append([])

        # Section 5: Educational Notes & Limitations
        ws_summary.append(["5. ACADEMIC NOTES & SPREADSHEET CAVEATS"])
        ws_summary.append(["• Relational Projection: Projected tuple duplicates were automatically removed in compliance with relational algebra."])
        ws_summary.append(["• Finite Sample Limit: Dependencies and keys confirmed in this report reflect the declared schema and verified algorithms, not mere finite row observations."])
        ws_summary.append(["• Source Preservation: The user's original uploaded workbook was not altered or overwritten."])

        # Style Section 1 summary items
        for r_i in range(5, 11):
            ws_summary.cell(row=r_i, column=1).font = bold_font
            ws_summary.cell(row=r_i, column=1).fill = accent_fill
            ws_summary.cell(row=r_i, column=2).font = regular_font

        ws_summary.column_dimensions["A"].width = 32
        ws_summary.column_dimensions["B"].width = 45
        ws_summary.column_dimensions["C"].width = 30
        ws_summary.column_dimensions["D"].width = 45
        ws_summary.column_dimensions["E"].width = 40

        # Save and return bytes
        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()
