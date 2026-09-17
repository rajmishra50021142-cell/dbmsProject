"""
Tests for Phase 10 Report Generation & Download Engine.
Validates TXT, DOCX, and PDF report builders and FastAPI endpoints.
"""

import io
import docx
import pytest
from fastapi.testclient import TestClient
from app.main import app
from app.reports.report_builder import (
    ReportSectionSelection,
    generate_text_report,
    generate_docx_report,
    generate_pdf_report,
    sanitize_filename,
)

client = TestClient(app)

MOCK_SCHEMA = {
    "name": "ENROLLMENT",
    "attributes": ["StudentID", "CourseID", "StudentName", "CourseName", "Grade"],
    "candidate_keys": [["StudentID", "CourseID"]],
    "functional_dependencies": [
        {"left": ["StudentID"], "right": ["StudentName"]},
        {"left": ["CourseID"], "right": ["CourseName"]},
        {"left": ["StudentID", "CourseID"], "right": ["Grade"]},
    ],
    "multivalued_dependencies": [],
}

MOCK_ANALYSIS_RESULT = {
    "relation_name": "ENROLLMENT",
    "attributes": ["StudentID", "CourseID", "StudentName", "CourseName", "Grade"],
    "candidate_keys": [["StudentID", "CourseID"]],
    "prime_attributes": ["StudentID", "CourseID"],
    "non_prime_attributes": ["StudentName", "CourseName", "Grade"],
    "summary_verdict": "Relation satisfies 1NF but violates 2NF due to partial dependencies.",
    "highest_confirmed_normal_form": "1NF",
    "nf1": {
        "status": "SATISFIED",
        "is_satisfied": True,
        "message": "All attributes are atomic and tuple identity is established.",
        "reasoning_steps": [{"explanation": "Cell values are atomic."}],
    },
    "nf2": {
        "status": "VIOLATED",
        "is_satisfied": False,
        "message": "Partial dependencies detected on proper subsets of candidate keys.",
        "partial_dependencies": [
            {
                "determinant": ["StudentID"],
                "dependent_attributes": ["StudentName"],
                "composite_key": ["StudentID", "CourseID"],
            },
            {
                "determinant": ["CourseID"],
                "dependent_attributes": ["CourseName"],
                "composite_key": ["StudentID", "CourseID"],
            },
        ],
        "reasoning_steps": [
            {"explanation": "StudentID is a proper subset of candidate key {StudentID, CourseID}."},
            {"explanation": "StudentID -> StudentName causes 2NF violation."},
        ],
        "decomposition_proposal": {
            "decomposed_relations": [
                {"name": "R1_Student", "attributes": ["StudentID", "StudentName"]},
                {"name": "R2_Course", "attributes": ["CourseID", "CourseName"]},
                {"name": "R3_Enrollment", "attributes": ["StudentID", "CourseID", "Grade"]},
            ],
            "reasoning": "Decompose into separate relations to isolate partial determinants.",
        },
    },
    "nf3": {
        "status": "BLOCKED_BY_PREREQUISITE",
        "is_satisfied": False,
        "message": "3NF analysis blocked because 2NF is violated.",
        "violations": [],
        "reasoning_steps": [],
    },
    "nf4": {
        "status": "BLOCKED_BY_PREREQUISITE",
        "is_satisfied": False,
        "message": "4NF analysis blocked because 3NF is not satisfied.",
        "violations": [],
        "reasoning_steps": [],
    },
    "lossless_join": {
        "is_lossless": True,
        "method": "TABLEAU_CHASE",
        "reasoning": "Tableau chase successfully equated all symbols to a distinguished row.",
    },
    "dependency_preservation": {
        "is_preserved": True,
        "reasoning": "All original dependencies are preserved within projections.",
    },
}


def test_sanitize_filename():
    assert sanitize_filename("ENROLLMENT", "pdf") == "normalization-report-ENROLLMENT.pdf"
    assert sanitize_filename("Student Record (2026)", "docx") == "normalization-report-Student_Record_2026.docx"
    assert sanitize_filename("", "txt") == "normalization-report-Relation.txt"


def test_generate_text_report():
    txt = generate_text_report(MOCK_SCHEMA, MOCK_ANALYSIS_RESULT)
    assert "DBMS NORMALIZATION LABORATORY" in txt
    assert "Dr. Swaminathan A" in txt
    assert "Assistant Professor" in txt
    assert "ENROLLMENT" in txt
    assert "StudentID, CourseID" in txt
    assert "VIOLATED" in txt
    assert "1NF" in txt
    assert "2NF" in txt


def test_generate_text_report_section_toggle():
    sections = ReportSectionSelection(user_inputs=False, figures=False)
    txt = generate_text_report(MOCK_SCHEMA, MOCK_ANALYSIS_RESULT, sections)
    assert "1. USER INPUT SPECIFICATION" not in txt
    assert "7. ARCHITECTURAL FIGURES" not in txt
    # Intermediate & Final results should still be present
    assert "2. INTERMEDIATE MATHEMATICAL DERIVATIONS" in txt
    assert "4. NORMALIZATION EVALUATION VERDICTS" in txt


def test_generate_docx_report():
    buf = generate_docx_report(MOCK_SCHEMA, MOCK_ANALYSIS_RESULT)
    assert isinstance(buf, io.BytesIO)
    assert buf.getvalue().startswith(b"PK")  # ZIP header for docx

    # Verify that python-docx can open it
    doc = docx.Document(buf)
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "DBMS Normalization Laboratory" in full_text
    cell_texts = [c.text for row in doc.tables[0].rows for c in row.cells]
    assert any("Dr. Swaminathan A" in ct for ct in cell_texts)


def test_generate_pdf_report():
    buf = generate_pdf_report(MOCK_SCHEMA, MOCK_ANALYSIS_RESULT)
    assert isinstance(buf, io.BytesIO)
    raw_pdf = buf.getvalue()
    assert raw_pdf.startswith(b"%PDF")
    assert len(raw_pdf) > 2000  # Non-trivial PDF size


def test_api_generate_txt():
    payload = {
        "format": "txt",
        "schema_definition": MOCK_SCHEMA,
        "analysis_result": MOCK_ANALYSIS_RESULT,
    }
    response = client.post("/api/v1/reports/generate", json=payload)
    assert response.status_code == 200
    assert "text/plain" in response.headers["content-type"]
    assert "attachment; filename=" in response.headers["content-disposition"]
    assert "normalization-report-ENROLLMENT.txt" in response.headers["content-disposition"]
    assert "Dr. Swaminathan A" in response.text


def test_api_generate_docx():
    payload = {
        "format": "docx",
        "schema_definition": MOCK_SCHEMA,
        "analysis_result": MOCK_ANALYSIS_RESULT,
    }
    response = client.post("/api/v1/reports/docx", json=payload)
    assert response.status_code == 200
    assert "wordprocessingml" in response.headers["content-type"]
    assert response.content.startswith(b"PK")


def test_api_generate_pdf():
    payload = {
        "format": "pdf",
        "schema_definition": MOCK_SCHEMA,
        "analysis_result": MOCK_ANALYSIS_RESULT,
    }
    response = client.post("/api/v1/reports/pdf", json=payload)
    assert response.status_code == 200
    assert "application/pdf" in response.headers["content-type"]
    assert response.content.startswith(b"%PDF")
